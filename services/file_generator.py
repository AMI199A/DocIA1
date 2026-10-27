import json
import os
import re
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from dotenv import load_dotenv
from jinja2 import Environment, FileSystemLoader, select_autoescape

load_dotenv()
STORAGE_PATH = os.getenv("STORAGE_PATH", "./archivos_generados")
TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates"
LOGO_PATH = Path(__file__).resolve().parent.parent / "static" / "logo.png"

# Asegurar que la carpeta exista
os.makedirs(STORAGE_PATH, exist_ok=True)


def sanitizar_texto_para_excel(texto: str) -> str:
    """Elimina caracteres de control ASCII que causan IllegalCharacterError en openpyxl."""
    if not texto:
        return ""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(texto))


def _normalizar_contenido(contenido: str | dict[str, Any]) -> dict[str, Any]:
    if isinstance(contenido, dict):
        data = dict(contenido)
    else:
        sections: list[dict[str, Any]] = []
        current_section: dict[str, Any] = {"title": "Contenido", "paragraphs": [], "bullets": []}
        for line in str(contenido).splitlines():
            text = line.strip()
            if not text:
                continue
            if text.startswith("#"):
                if current_section["paragraphs"] or current_section["bullets"]:
                    sections.append(current_section)
                current_section = {"title": text.lstrip("# ").strip(), "paragraphs": [], "bullets": []}
            elif text.startswith(("- ", "* ", "• ")):
                current_section["bullets"].append(text[2:].strip())
            else:
                current_section["paragraphs"].append(text)
        if current_section["paragraphs"] or current_section["bullets"]:
            sections.append(current_section)
        data = {"title": "Reporte DocIA PRO", "sections": sections}

    data.setdefault("title", "Reporte DocIA PRO")
    data.setdefault("subtitle", "")
    data.setdefault("summary", "")
    data.setdefault("sections", [])
    data.setdefault("slides", [])
    data.setdefault("qa", [])
    data.setdefault("key_points", [])
    return data


def contenido_a_markdown(contenido: str | dict[str, Any]) -> str:
    data = _normalizar_contenido(contenido)
    lines = [f"# {data['title']}"]
    if data["subtitle"]:
        lines.extend(["", str(data["subtitle"])])
    if data["summary"]:
        lines.extend(["", "## Resumen ejecutivo", "", str(data["summary"])])
    for section in data["sections"]:
        lines.extend(["", f"## {section.get('title', 'Sección')}"])
        lines.append("")
        lines.extend(str(paragraph) for paragraph in section.get("paragraphs", []) if paragraph)
        lines.extend(f"- {bullet}" for bullet in section.get("bullets", []) if bullet)
    if data["slides"]:
        lines.extend(["", "## Diapositivas"])
        for slide in data["slides"]:
            lines.extend(["", f"### {slide.get('title', 'Diapositiva')}"])
            lines.extend(f"- {bullet}" for bullet in slide.get("bullets", []) if bullet)
    if data["qa"]:
        lines.extend(["", "## Preguntas y respuestas"])
        for item in data["qa"]:
            lines.extend(["", f"**Pregunta:** {item.get('question', '')}", f"**Respuesta:** {item.get('answer', '')}"])
    if data["key_points"]:
        lines.extend(["", "## Puntos clave"])
        lines.extend(f"- {point}" for point in data["key_points"] if point)
    return "\n".join(lines).strip()


def _render_html(contenido: str | dict[str, Any]) -> str:
    environment = Environment(
        loader=FileSystemLoader(str(TEMPLATE_DIR)),
        autoescape=select_autoescape(("html", "xml")),
    )
    template = environment.get_template("reporte.html")
    return template.render(report=_normalizar_contenido(contenido), generated_at=datetime.now().strftime("%d/%m/%Y %H:%M"))


def crear_pdf(contenido: str | dict[str, Any], nombre_archivo: str) -> str:
    ruta = os.path.join(STORAGE_PATH, f"{nombre_archivo}.pdf")
    from weasyprint import HTML

    HTML(string=_render_html(contenido), base_url=str(TEMPLATE_DIR)).write_pdf(ruta)
    return ruta


def crear_docx(contenido: str | dict[str, Any], nombre_archivo: str) -> str:
    ruta = os.path.join(STORAGE_PATH, f"{nombre_archivo}.docx")
    data = _normalizar_contenido(contenido)
    template_path = os.getenv("DOCX_TEMPLATE_PATH", "")
    doc = Document(template_path) if template_path and os.path.isfile(template_path) else Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(45, 55, 72)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[style_name].font.name = "Aptos Display"
        doc.styles[style_name].font.color.rgb = RGBColor(22, 78, 99)

    doc.add_heading(data["title"], 0)
    if data["subtitle"]:
        subtitle = doc.add_paragraph(data["subtitle"])
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if data["summary"]:
        doc.add_heading("Resumen ejecutivo", level=1)
        doc.add_paragraph(data["summary"])
    for item in data["sections"]:
        doc.add_heading(item.get("title", "Sección"), level=1)
        for paragraph in item.get("paragraphs", []):
            doc.add_paragraph(paragraph)
        for bullet in item.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")
    if data["slides"]:
        doc.add_heading("Contenido de la presentación", level=1)
        for slide in data["slides"]:
            doc.add_heading(slide.get("title", "Diapositiva"), level=2)
            for bullet in slide.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")
    if data["qa"]:
        doc.add_heading("Preguntas y respuestas", level=1)
        for item in data["qa"]:
            question = doc.add_paragraph()
            question.add_run(f"Pregunta: {item.get('question', '')}").bold = True
            doc.add_paragraph(f"Respuesta: {item.get('answer', '')}")
    if data["key_points"]:
        doc.add_heading("Puntos clave", level=1)
        for point in data["key_points"]:
            doc.add_paragraph(point, style="List Bullet")
    doc.save(ruta)
    return ruta


def _convertir_a_spec_presentacion(data: dict[str, Any]) -> dict[str, Any]:
    title = data.get("title") or "Reporte DocIA PRO"
    subtitle = data.get("subtitle") or data.get("summary") or "Inteligencia Documental Avanzada"
    summary = data.get("summary") or ""

    slides: list[dict[str, Any]] = [
        {
            "layout": "cover",
            "title": title,
            "subtitle": subtitle,
            "notes": f"Presentación generada por DocIA PRO. {summary}".strip(),
        }
    ]

    slides_data = list(data.get("slides") or [])
    if not slides_data:
        for section in data.get("sections") or []:
            bullets = [str(b) for b in section.get("bullets", [])] + [str(p) for p in section.get("paragraphs", [])]
            if bullets:
                slides_data.append({
                    "title": section.get("title", "Sección"),
                    "bullets": bullets[:6],
                })
    if not slides_data and data.get("key_points"):
        slides_data.append({
            "title": "Puntos Clave",
            "bullets": [str(p) for p in data.get("key_points", [])][:6],
        })
    if data.get("qa"):
        for index, item in enumerate(data.get("qa", []), start=1):
            slides_data.append({
                "title": f"Pregunta {index}: {item.get('question', '')}"[:100],
                "bullets": [str(item.get("answer", ""))],
            })
    if not slides_data:
        slides_data.append({
            "title": "Resumen Ejecutivo",
            "bullets": [summary or "Sin contenido disponible."],
        })

    for s in slides_data:
        raw_bullets = s.get("bullets", [])
        if isinstance(raw_bullets, str):
            raw_bullets = [raw_bullets]
        slides.append({
            "layout": "content",
            "title": s.get("title", "Diapositiva"),
            "bullets": [str(b).strip() for b in raw_bullets if str(b).strip()][:7],
            "notes": s.get("notes", ""),
        })

    return {
        "title": title,
        "slides": slides,
    }


def _crear_pptx_con_node(spec: dict[str, Any], ruta_salida: str) -> bool:
    cli_path = Path(__file__).resolve().parent.parent / "src" / "presentation" / "cli.js"
    if not cli_path.exists():
        return False

    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".json", delete=False) as f:
        json.dump(spec, f, ensure_ascii=False)
        temp_spec_path = f.name

    try:
        res = subprocess.run(
            ["node", str(cli_path), temp_spec_path, ruta_salida],
            capture_output=True,
            text=True,
            timeout=30,
        )
        return res.returncode == 0 and os.path.isfile(ruta_salida) and os.path.getsize(ruta_salida) > 0
    except Exception:
        return False
    finally:
        if os.path.exists(temp_spec_path):
            try:
                os.remove(temp_spec_path)
            except OSError:
                pass


def _crear_pptx_con_python(spec: dict[str, Any], ruta_salida: str) -> str:
    from pptx import Presentation
    from pptx.util import Inches as PInches, Pt as PPt
    from pptx.dml.color import RGBColor as PRGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = PInches(10)
    prs.slide_height = PInches(5.625)
    blank_layout = prs.slide_layouts[6]

    logo_path = LOGO_PATH
    has_logo = logo_path.is_file()
    logo_str = str(logo_path)

    # Colores del tema DocIA PRO (según generador_presentaciones.txt)
    color_primary = PRGBColor(0x1F, 0x3A, 0x5F)
    color_accent = PRGBColor(0xF2, 0xA5, 0x41)
    color_text = PRGBColor(0x22, 0x22, 0x22)
    color_light = PRGBColor(0xFF, 0xFF, 0xFF)
    color_muted = PRGBColor(0x6B, 0x72, 0x80)

    for slide_idx, s in enumerate(spec.get("slides", []), start=1):
        is_cover = s.get("layout") == "cover" or slide_idx == 1
        slide = prs.slides.add_slide(blank_layout)

        if is_cover:
            # Fondo primario
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PInches(10), PInches(5.625))
            bg.fill.solid()
            bg.fill.fore_color.rgb = color_primary
            bg.line.fill.background()

            # Logo DocIA PRO en portada
            if has_logo:
                slide.shapes.add_picture(logo_str, PInches(0.7), PInches(0.5), width=PInches(1.1))

            # Título Cover
            tb_title = slide.shapes.add_textbox(PInches(0.7), PInches(1.8), PInches(8.6), PInches(1.4))
            tf_title = tb_title.text_frame
            tf_title.word_wrap = True
            tf_title.vertical_anchor = MSO_ANCHOR.BOTTOM
            p_title = tf_title.paragraphs[0]
            p_title.text = s.get("title", "")
            p_title.font.name = "Georgia"
            p_title.font.size = PPt(40)
            p_title.font.bold = True
            p_title.font.color.rgb = color_light

            # Subtítulo Cover
            subtitle_text = s.get("subtitle", "")
            if subtitle_text:
                tb_sub = slide.shapes.add_textbox(PInches(0.7), PInches(3.35), PInches(8.6), PInches(0.8))
                tf_sub = tb_sub.text_frame
                tf_sub.word_wrap = True
                tf_sub.vertical_anchor = MSO_ANCHOR.TOP
                p_sub = tf_sub.paragraphs[0]
                p_sub.text = str(subtitle_text)
                p_sub.font.name = "Arial"
                p_sub.font.size = PPt(18)
                p_sub.font.color.rgb = color_accent

        else:
            # Diapositiva de contenido
            # Franja superior accent
            top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PInches(10), PInches(0.12))
            top_bar.fill.solid()
            top_bar.fill.fore_color.rgb = color_accent
            top_bar.line.fill.background()

            # Logo en la esquina superior derecha
            if has_logo:
                slide.shapes.add_picture(logo_str, PInches(8.95), PInches(0.35), width=PInches(0.55))

            # Título Diapositiva
            tb_title = slide.shapes.add_textbox(PInches(0.6), PInches(0.4), PInches(8.2), PInches(0.9))
            tf_title = tb_title.text_frame
            tf_title.word_wrap = True
            tf_title.vertical_anchor = MSO_ANCHOR.MIDDLE
            p_title = tf_title.paragraphs[0]
            p_title.text = s.get("title", "Contenido")
            p_title.font.name = "Georgia"
            p_title.font.size = PPt(28)
            p_title.font.bold = True
            p_title.font.color.rgb = color_primary

            # Viñetas / Cuerpo
            tb_body = slide.shapes.add_textbox(PInches(0.6), PInches(1.5), PInches(8.8), PInches(3.5))
            tf_body = tb_body.text_frame
            tf_body.word_wrap = True
            bullets = s.get("bullets", [])
            for b_idx, bullet in enumerate(bullets):
                p_b = tf_body.paragraphs[0] if b_idx == 0 else tf_body.add_paragraph()
                p_b.text = f"•  {bullet}"
                p_b.font.name = "Arial"
                p_b.font.size = PPt(18)
                p_b.font.color.rgb = color_text
                p_b.space_after = PPt(8)

            # Pie de página: Logo miniatura + Marca "DocIA PRO"
            if has_logo:
                slide.shapes.add_picture(logo_str, PInches(0.6), PInches(5.18), width=PInches(0.28))

            tb_foot = slide.shapes.add_textbox(PInches(0.95), PInches(5.15), PInches(5.5), PInches(0.3))
            tf_foot = tb_foot.text_frame
            p_foot = tf_foot.paragraphs[0]
            p_foot.text = "DocIA PRO"
            p_foot.font.name = "Arial"
            p_foot.font.size = PPt(10)
            p_foot.font.bold = True
            p_foot.font.color.rgb = color_muted

            # Número de diapositiva
            tb_num = slide.shapes.add_textbox(PInches(8.8), PInches(5.15), PInches(0.6), PInches(0.3))
            tf_num = tb_num.text_frame
            p_num = tf_num.paragraphs[0]
            p_num.text = str(slide_idx)
            p_num.alignment = PP_ALIGN.RIGHT
            p_num.font.name = "Arial"
            p_num.font.size = PPt(10)
            p_num.font.color.rgb = color_muted

        # Notas de orador si existen
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s.get("notes"))

    prs.save(ruta_salida)
    return ruta_salida


def crear_pptx(contenido: str | dict[str, Any], nombre_archivo: str) -> str:
    ruta = os.path.join(STORAGE_PATH, f"{nombre_archivo}.pptx")
    data = _normalizar_contenido(contenido)
    spec = _convertir_a_spec_presentacion(data)

    # Intento 1: Generador nativo Node.js con pptxgenjs (según especificación)
    if _crear_pptx_con_node(spec, ruta):
        return ruta

    # Intento 2 / Respaldo robusto: Generador Python-pptx con diseño 16:9 y tema idéntico
    return _crear_pptx_con_python(spec, ruta)


def crear_md(contenido: str | dict[str, Any], nombre_archivo: str) -> str:
    ruta = os.path.join(STORAGE_PATH, f"{nombre_archivo}.md")
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(contenido_a_markdown(contenido))
    return ruta


def crear_txt(contenido: str | dict[str, Any], nombre_archivo: str) -> str:
    ruta = os.path.join(STORAGE_PATH, f"{nombre_archivo}.txt")
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(contenido_a_markdown(contenido).replace("#", "").replace("**", ""))
    return ruta


def crear_html(contenido: str | dict[str, Any], nombre_archivo: str) -> str:
    ruta = os.path.join(STORAGE_PATH, f"{nombre_archivo}.html")
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(_render_html(contenido))
    return ruta


def crear_xlsx(contenido: str | dict[str, Any], nombre_archivo: str) -> str:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    data = _normalizar_contenido(contenido)
    ruta = os.path.join(STORAGE_PATH, f"{nombre_archivo}.xlsx")
    wb = openpyxl.Workbook()
    header_fill = PatternFill(fill_type="solid", fgColor="164E63")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(vertical="top", wrap_text=True)

    def formatear_hoja(ws, column_widths: dict[str, int]) -> None:
        for column, width in column_widths.items():
            ws.column_dimensions[column].width = width
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.value = sanitizar_texto_para_excel(cell.value)
                    cell.data_type = "s"
                    cell.alignment = wrap
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28

    ws = wb.active
    ws.title = "Reporte"
    ws.append(["Contenido", "Detalle"])
    ws.append(["Título", data["title"]])
    if data["subtitle"]:
        ws.append(["Subtítulo", data["subtitle"]])
    if data["summary"]:
        ws.append(["Resumen ejecutivo", data["summary"]])
    for section in data["sections"]:
        ws.append([section.get("title", "Sección"), ""])
        for paragraph in section.get("paragraphs", []):
            ws.append(["", paragraph])
        for bullet in section.get("bullets", []):
            ws.append(["", f"• {bullet}"])
    for point in data["key_points"]:
        ws.append(["Punto clave", f"• {point}"])
    formatear_hoja(ws, {"A": 30, "B": 100})

    if data["qa"]:
        ws = wb.create_sheet("Cuestionario")
        ws.append(["Pregunta", "Respuesta"])
        for item in data["qa"]:
            ws.append([
                sanitizar_texto_para_excel(item.get("question", "")),
                sanitizar_texto_para_excel(item.get("answer", "")),
            ])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        formatear_hoja(ws, {"A": 48, "B": 90})

    if data["slides"]:
        ws = wb.create_sheet("Diapositivas")
        ws.append(["Diapositiva", "Puntos"])
        for slide in data["slides"]:
            ws.append([
                slide.get("title", "Diapositiva"),
                "\n".join(f"• {bullet}" for bullet in slide.get("bullets", [])),
            ])
        formatear_hoja(ws, {"A": 36, "B": 100})

    wb.save(ruta)
    return ruta


def generar_archivo_salida(contenido: str | dict[str, Any], formato: str, nombre_base: str) -> str:
    fmt = formato.lower().strip().replace(".", "")
    if fmt == "pdf":
        return crear_pdf(contenido, nombre_base)
    elif fmt == "docx":
        return crear_docx(contenido, nombre_base)
    elif fmt == "pptx":
        return crear_pptx(contenido, nombre_base)
    elif fmt == "md":
        return crear_md(contenido, nombre_base)
    elif fmt == "txt":
        return crear_txt(contenido, nombre_base)
    elif fmt == "xlsx":
        return crear_xlsx(contenido, nombre_base)
    elif fmt == "html":
        return crear_html(contenido, nombre_base)
    raise ValueError(f"Formato de salida no soportado: {formato}")
