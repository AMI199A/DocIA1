import io
import os
import logging
import traceback
from typing import Tuple

import pypdf
import docx

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("file_extractor")

ALLOWED_EXTENSIONS = {"pdf", "docx", "doc", "txt", "md", "rtf"}
ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/x-pdf",
    "application/acrobat",
    "applications/vnd.pdf",
    "text/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "text/plain",
    "text/markdown",
    "application/octet-stream"
}

def extraer_texto_pdf(file_bytes: bytes, filename: str = "") -> str:
    """
    Extrae texto de un archivo PDF usando pypdf con manejo de páginas protegidas o dañadas.
    """
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes), strict=False)
        
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as dec_err:
                logger.warning(f"PDF protegido o encriptado ({filename}): {dec_err}")
                raise ValueError("El archivo PDF está protegido con contraseña o encriptado.")
        
        num_paginas = len(reader.pages)
        logger.info(f"Procesando PDF '{filename}' con {num_paginas} páginas.")
        
        paginas_texto = []
        for i, page in enumerate(reader.pages):
            try:
                texto_pagina = page.extract_text()
                if texto_pagina:
                    paginas_texto.append(texto_pagina.strip())
            except Exception as page_err:
                logger.warning(f"Advertencia al extraer página {i+1} de '{filename}': {page_err}")
                continue
                
        texto_final = "\n\n".join(paginas_texto).strip()
        if not texto_final:
            raise ValueError("No se pudo extraer texto legible del PDF. Puede tratarse de un documento escaneado como imagen.")
            
        logger.info(f"Extracción exitosa de '{filename}': {len(texto_final)} caracteres extraídos.")
        return texto_final
    except ValueError:
        raise
    except Exception as e:
        logger.error(f"Error inesperado procesando PDF '{filename}':\n{traceback.format_exc()}")
        raise ValueError(f"Error al procesar el archivo PDF: {str(e)}")

def extraer_texto_docx(file_bytes: bytes, filename: str = "") -> str:
    """
    Extrae texto de un archivo DOCX incluyendo párrafos y tablas.
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        parrafos = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Extraer también celdas de tablas si existen
        for table in doc.tables:
            for row in table.rows:
                fila_texto = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if fila_texto:
                    parrafos.append(fila_texto)
                    
        texto_final = "\n\n".join(parrafos).strip()
        if not texto_final:
            raise ValueError("El documento Word no contiene texto extraíble.")
            
        logger.info(f"Extracción exitosa de DOCX '{filename}': {len(texto_final)} caracteres.")
        return texto_final
    except Exception as e:
        logger.error(f"Error procesando DOCX '{filename}':\n{traceback.format_exc()}")
        raise ValueError(f"Error al procesar el archivo Word (.docx): {str(e)}")

def extraer_texto_txt(file_bytes: bytes, filename: str = "") -> str:
    """
    Decodifica archivo de texto plano con soporte para múltiples codificaciones (UTF-8, Latin-1, CP1252).
    """
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
    for enc in encodings:
        try:
            texto = file_bytes.decode(enc).strip()
            if texto:
                logger.info(f"Texto TXT '{filename}' decodificado exitosamente con {enc}.")
                return texto
        except UnicodeDecodeError:
            continue
    raise ValueError("No se pudo decodificar el archivo de texto plano.")

def extraer_texto_de_archivo(file_bytes: bytes, filename: str, mime_type: str = "") -> str:
    """
    Función unificada para extraer texto de archivos PDF, DOCX o TXT con detección inteligente.
    """
    if not filename:
        filename = "documento.pdf"
        
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    logger.info(f"Solicitud de extracción para archivo: '{filename}', extensión: '{ext}', MIME: '{mime_type}'")
    
    # 1. Detección PDF (por extensión, MIME o cabecera %PDF)
    if ext == "pdf" or (mime_type and "pdf" in mime_type.lower()) or file_bytes.startswith(b'%PDF'):
        return extraer_texto_pdf(file_bytes, filename)
    
    # 2. Detección Word DOCX (por extensión, MIME o cabecera ZIP/PK)
    elif ext in ("docx", "doc") or (mime_type and "word" in mime_type.lower()) or file_bytes.startswith(b'PK\x03\x04'):
        try:
            return extraer_texto_docx(file_bytes, filename)
        except Exception:
            return extraer_texto_txt(file_bytes, filename)
            
    # 3. Detección Texto plano / Markdown
    elif ext in ("txt", "md", "rtf", "csv", "json") or (mime_type and "text" in mime_type.lower()):
        return extraer_texto_txt(file_bytes, filename)
        
    # 4. Fallback permisivo inteligente
    else:
        try:
            if file_bytes.startswith(b'%PDF'):
                return extraer_texto_pdf(file_bytes, filename)
            return extraer_texto_txt(file_bytes, filename)
        except Exception:
            logger.error(f"Formato no soportado: '{ext}' (MIME: {mime_type}) en archivo '{filename}'")
            raise ValueError(f"Formato no soportado (.{ext}). Los formatos permitidos son: .pdf, .docx, .txt")
